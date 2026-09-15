import os
import re
import math
import shutil
import urllib.request
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Cm as PptCm, Pt as PptPt


BASE = Path(r"D:\Desktop")
SRC_DOC = BASE / "文档.docx"
TEMPLATE = BASE / "模板ppt.pptx"
IMG_DIR = BASE / "新建文件夹" / "fuzzy_idm_results_v2"
PDF_DIR = BASE / "文献PDF文件夹"
OUT_PPT = BASE / "阶段学习组会汇报.pptx"
OUT_STAGE3 = BASE / "阶段三_多模态模糊神经微分方程构思.docx"
OUT_LIT = BASE / "文献清单.xlsx"
OUT_REVIEW = BASE / "阶段三文献调研小结.docx"

FONT_CN = "等线"
FONT_EN = "Times New Roman"
BLUE = PptRGB(31, 78, 121)
LIGHT_BLUE = PptRGB(221, 235, 247)
ORANGE = PptRGB(237, 125, 49)
GREEN = PptRGB(112, 173, 71)
GRAY = PptRGB(89, 89, 89)


def clear_slides(prs):
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def set_run_font(run, size=None, bold=False, color=None, italic=False):
    run.font.name = FONT_CN
    run._r.rPr.set(qn("a:ea"), FONT_CN)
    run._r.rPr.set(qn("a:latin"), FONT_EN)
    if size:
        run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, text="", size=20, color=PptRGB(0, 0, 0),
                bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
                fill=None, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    box = slide.shapes.add_shape(shape_type, PptCm(x), PptCm(y), PptCm(w), PptCm(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = PptPt(1)
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 1.0, 0.45, 31.8, 1.0, title, size=26, color=BLUE, bold=True)
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, PptCm(1.0), PptCm(1.55), PptCm(31.8), PptCm(0.04)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLUE
    slide.shapes[-1].line.fill.background()
    if subtitle:
        add_textbox(slide, 1.05, 1.65, 31, 0.55, subtitle, size=12, color=GRAY)


def add_footer(slide, n):
    add_textbox(slide, 28.8, 18.3, 3.8, 0.45, f"{n:02d}/16", size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def add_bullets(slide, x, y, w, h, items, size=16, color=PptRGB(30, 30, 30),
                bullet=True, level=0, spacing=4):
    box = slide.shapes.add_textbox(PptCm(x), PptCm(y), PptCm(w), PptCm(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.space_after = PptPt(spacing)
        p.alignment = PP_ALIGN.LEFT
        if bullet:
            p.text = f"• {item}"
            for run in p.runs:
                set_run_font(run, size=size, color=color)
        else:
            run = p.add_run()
            run.text = item
            set_run_font(run, size=size, color=color)
    return box


def add_picture_fit(slide, image_path, x, y, w, h, caption=None):
    image_path = Path(image_path)
    with Image.open(image_path) as im:
        iw, ih = im.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        draw_w = w
        draw_h = w / img_ratio
        draw_x = x
        draw_y = y + (h - draw_h) / 2
    else:
        draw_h = h
        draw_w = h * img_ratio
        draw_x = x + (w - draw_w) / 2
        draw_y = y
    pic = slide.shapes.add_picture(str(image_path), PptCm(draw_x), PptCm(draw_y), PptCm(draw_w), PptCm(draw_h))
    if caption:
        add_textbox(slide, x, y + h + 0.05, w, 0.45, caption, size=10, color=GRAY, align=PP_ALIGN.CENTER)
    return pic


def add_flow_box(slide, x, y, w, h, text, fill=LIGHT_BLUE, color=BLUE, size=13):
    return add_textbox(slide, x, y, w, h, text, size=size, color=color, bold=True,
                       align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, fill=fill, line=BLUE, radius=True)


def add_arrow(slide, x1, y1, x2, y2, color=BLUE):
    line = slide.shapes.add_connector(1, PptCm(x1), PptCm(y1), PptCm(x2), PptCm(y2))
    line.line.color.rgb = color
    line.line.width = PptPt(1.5)
    try:
        line.line.end_arrowhead = True
    except Exception:
        pass
    return line


def table(slide, x, y, w, h, rows, col_widths=None, font_size=11, header_fill=LIGHT_BLUE):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), PptCm(x), PptCm(y), PptCm(w), PptCm(h))
    tbl = shp.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = PptCm(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.margin_left = PptCm(0.08)
            cell.margin_right = PptCm(0.08)
            cell.margin_top = PptCm(0.03)
            cell.margin_bottom = PptCm(0.03)
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if r == 0 else PP_ALIGN.LEFT
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(r == 0), color=PptRGB(0, 0, 0))
    return shp


def create_ppt():
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)
    blank = prs.slide_layouts[0]

    # 1
    s = prs.slides.add_slide(blank)
    add_textbox(s, 2.1, 4.2, 29.5, 1.35, "模糊微分方程与 Fuzzy-IDM 阶段学习汇报", 31, BLUE, True, PP_ALIGN.CENTER)
    add_textbox(s, 2.1, 5.8, 29.5, 0.8, "阶段一理论学习与阶段二仿真分析", 20, GRAY, False, PP_ALIGN.CENTER)
    add_textbox(s, 2.1, 8.1, 29.5, 0.75, "汇报人：硕士新生    日期：2026年9月13日", 15, PptRGB(60, 60, 60), False, PP_ALIGN.CENTER)
    add_textbox(s, 7.6, 11.8, 18.5, 0.8, "第一次组会汇报：以学习理解、初步复现和后续构思为主", 14, BLUE, False, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_footer(s, 1)

    # 2
    s = prs.slides.add_slide(blank); add_title(s, "研究背景与任务来源"); add_footer(s, 2)
    add_bullets(s, 1.2, 2.4, 15.2, 8.0, [
        "微观交通模型需要描述车辆跟驰中的速度、间距、相对速度和加速度演化。",
        "传统 IDM 参数通常取确定值，但真实驾驶行为和感知输入具有明显不确定性。",
        "相机、激光雷达等交通感知数据会受到遮挡、光照、天气、点云稀疏和检测框边界不清等因素影响。",
        "这些不确定性不一定适合直接假设为概率分布，可先用模糊数与 α-截集进行区间化表达。"
    ], size=15)
    add_flow_box(s, 18.0, 3.1, 4.0, 1.35, "感知误差")
    add_flow_box(s, 23.0, 3.1, 4.0, 1.35, "状态不确定")
    add_flow_box(s, 28.0, 3.1, 4.0, 1.35, "跟驰预测区间")
    add_arrow(s, 22.0, 3.78, 23.0, 3.78); add_arrow(s, 27.0, 3.78, 28.0, 3.78)
    add_textbox(s, 18.0, 5.3, 14.0, 4.0,
                "本阶段汇报的定位：\n\n从模糊微分方程的基础概念出发，理解如何把 IDM 的关键驾驶参数模糊化，并通过 Python 仿真观察不同交通场景下的不确定性传播。", 15, PptRGB(30,30,30),
                fill=PptRGB(242, 246, 250), line=PptRGB(200, 210, 220), radius=True)

    # 3
    s = prs.slides.add_slide(blank); add_title(s, "阶段任务总体安排"); add_footer(s, 3)
    xs = [2.2, 12.4, 22.6]
    titles = ["阶段一\n理论基础", "阶段二\n模型与仿真", "阶段三\n后续构思"]
    bodies = ["模糊数\nα-截集\nH差 / gH差\nFDE 数值求解", "IDM 公式理解\n三参数模糊化\nα-截集采样\n三类场景仿真", "多模态输入\n模糊神经网络\nFuzzy Neural ODE\n轨迹预测区间"]
    for i, x in enumerate(xs):
        add_flow_box(s, x, 3.0, 7.8, 1.35, titles[i], fill=[LIGHT_BLUE, PptRGB(252, 228, 214), PptRGB(226, 239, 218)][i], color=BLUE)
        add_textbox(s, x, 4.65, 7.8, 5.0, bodies[i], 15, PptRGB(45,45,45), align=PP_ALIGN.CENTER, fill=PptRGB(250,250,250), line=PptRGB(210,210,210), radius=True)
    add_arrow(s, 10.0, 3.65, 12.4, 3.65); add_arrow(s, 20.2, 3.65, 22.6, 3.65)
    add_bullets(s, 2.2, 11.1, 29.0, 2.1, [
        "当前已完成：阶段一理论学习、阶段二 Fuzzy-IDM 推导与初步仿真。",
        "本次重点汇报：我对模型的理解、仿真结果能说明什么，以及阶段三可如何继续推进。"
    ], size=15)

    # 4
    s = prs.slides.add_slide(blank); add_title(s, "模糊数与 α-截集"); add_footer(s, 4)
    add_picture_fit(s, BASE / "几何理解.png", 1.2, 2.25, 13.5, 8.0, "三角模糊数与 α-截集的几何理解")
    add_textbox(s, 15.7, 2.25, 16.5, 2.1, "三角模糊数：Ã = (l, m, r)\n隶属度越高，取值越接近核心值 m。", 17, PptRGB(30,30,30), fill=PptRGB(250,250,250), line=PptRGB(210,210,210), radius=True)
    add_textbox(s, 15.7, 5.0, 16.5, 2.0, "[Ã]α = [l + α(m-l),  r - α(r-m)]\n0 ≤ α ≤ 1", 20, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_bullets(s, 15.9, 7.8, 15.8, 4.0, [
        "α 越小，保留的不确定区间越宽。",
        "α=1 对应三角模糊数的核心值。",
        "不同 α 水平形成嵌套区间，是把 FDE 转换为常微分方程族的基础。"
    ], size=15)

    # 5
    s = prs.slides.add_slide(blank); add_title(s, "模糊微分方程基本思想"); add_footer(s, 5)
    add_textbox(s, 1.2, 2.45, 30.8, 1.5,
                "模糊状态随时间变化时，导数不能简单照搬实数减法，因为两个模糊数之间不一定存在普通意义下的“差”。", 18, PptRGB(30,30,30),
                fill=PptRGB(242,246,250), line=PptRGB(205,215,225), radius=True)
    table(s, 1.4, 4.6, 15.0, 6.2, [
        ["概念", "作用", "我的理解"],
        ["H差", "定义模糊数 A 与 B 的差", "适合模糊宽度扩张的情形"],
        ["gH差", "放宽差的存在条件", "可处理宽度收缩或方向变化"],
        ["H/gH导数", "定义模糊值函数导数", "让 FDE 能够进入数值求解"]
    ], col_widths=[3.2, 5.3, 6.5], font_size=11)
    add_textbox(s, 18.0, 4.6, 13.5, 2.2, "一般形式\nx̃′(t)=f(t,x̃(t)),  x̃(t₀)=x̃₀", 19, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_bullets(s, 18.2, 7.35, 13.0, 3.6, [
        "强解要求模糊状态本身满足导数条件。",
        "弱解可通过积分形式理解，适用范围更宽。",
        "实际计算常通过 α-截集转为上下端点 ODE。"
    ], size=15)

    # 6
    s = prs.slides.add_slide(blank); add_title(s, "简单 FDE 算例与数值方法"); add_footer(s, 6)
    add_textbox(s, 1.2, 2.15, 30.8, 1.05,
                "x̃′(t)=0.5x̃(t),  x̃(0)=(0.8, 1.0, 1.2)    在 α-截集下分别求上下端点。", 19, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_picture_fit(s, BASE / "解析解.png", 1.1, 3.7, 9.9, 7.1, "解析解/模糊带")
    add_picture_fit(s, BASE / "三法对比.png", 11.6, 3.7, 9.9, 7.1, "Euler、RK4 与解析解对比")
    add_picture_fit(s, BASE / "绝对误差.png", 22.1, 3.7, 9.9, 7.1, "数值误差对比")
    add_bullets(s, 1.3, 11.7, 30.5, 2.0, [
        "Euler 方法直观、实现简单，但长时间计算误差累积更明显。",
        "RK4 在相同步长下精度更高，因此后续复杂动力学仿真更适合采用较小步长和稳定数值积分。"
    ], size=14)

    # 7
    s = prs.slides.add_slide(blank); add_title(s, "经典 IDM 模型回顾"); add_footer(s, 7)
    add_textbox(s, 1.15, 2.35, 31.0, 2.0,
                "aₙ(t)=aₘₐₓ⁽ⁿ⁾[1-(Vₙ(t)/Ṽₙ(t))ᵝ-(S̃ₙ(t)/Sₙ(t))²]\nS̃ₙ(t)=Sⱼₐₘ⁽ⁿ⁾+Vₙ(t)T̃ₙ(t)+Vₙ(t)ΔVₙ(t)/(2√(aₘₐₓ⁽ⁿ⁾a꜀ₒₘf⁽ⁿ⁾))",
                17, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    table(s, 1.2, 5.25, 30.8, 5.1, [
        ["符号", "含义", "本阶段理解重点"],
        ["aₙ(t)", "第 n 辆车加速度", "IDM 的核心输出"],
        ["aₘₐₓ⁽ⁿ⁾ / a꜀ₒₘf⁽ⁿ⁾", "最大加速度 / 舒适减速度", "反映驾驶员加减速能力差异"],
        ["Vₙ(t) / Ṽₙ(t)", "实际速度 / 期望速度", "注意 Ṽₙ 是期望量标记"],
        ["Sₙ(t) / S̃ₙ(t)", "实际净间距 / 期望动态间距", "注意 S̃ₙ 是期望量标记"],
        ["T̃ₙ(t), Sⱼₐₘ⁽ⁿ⁾, β", "期望车头时距、静止间距、加速度指数", "控制跟驰安全性与响应强度"]
    ], col_widths=[5.4, 10.3, 15.1], font_size=10.5)
    add_textbox(s, 1.3, 11.2, 30.5, 1.8,
                "符号说明：原 IDM 中 Ṽₙ、S̃ₙ、T̃ₙ 上方波浪线主要表示“期望量”，不直接等同于本文的模糊符号；本阶段的模糊性通过 aₘₐₓ、a꜀ₒₘf、T 的 α-截集表达。", 14, PptRGB(90,55,0),
                fill=PptRGB(255,242,204), line=PptRGB(240,190,100), radius=True)

    # 8
    s = prs.slides.add_slide(blank); add_title(s, "三参数 Fuzzy-IDM 建模思路"); add_footer(s, 8)
    add_bullets(s, 1.2, 2.35, 30.0, 1.8, [
        "主动模糊化的只有三个参数：aₘₐₓ、a꜀ₒₘf、T；其余状态由动力学传播后成为模糊输出。"
    ], size=16)
    labels = ["模糊参数\naₘₐₓ, a꜀ₒₘf, T", "期望动态间距\nS̃ₙ(t)", "加速度\naₙ(t)", "速度\nVₙ(t)", "位置/间距\nXₙ(t), Sₙ(t)", "反馈到下一步"]
    xs = [1.2, 6.5, 11.8, 17.1, 22.4, 27.7]
    for x, lab in zip(xs, labels):
        add_flow_box(s, x, 5.1, 4.3, 1.65, lab, fill=PptRGB(242,246,250), size=12)
    for i in range(len(xs)-1):
        add_arrow(s, xs[i]+4.3, 5.92, xs[i+1], 5.92)
    add_textbox(s, 4.5, 8.5, 24.5, 2.0,
                "参数不确定性并不是只停留在参数层，而是随着 IDM 的非线性耦合关系传到速度、位置、实际净间距、相对速度和加速度。", 17, PptRGB(30,30,30),
                align=PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)

    # 9
    s = prs.slides.add_slide(blank); add_title(s, "α-截集下的数值求解流程"); add_footer(s, 9)
    add_textbox(s, 1.3, 2.35, 30.5, 1.35,
                "由于 IDM 是非线性耦合模型，不能简单把参数上下界代入后就认为得到状态上下界。", 18, PptRGB(90,55,0),
                align=PP_ALIGN.CENTER, fill=PptRGB(255,242,204), line=PptRGB(240,190,100), radius=True)
    steps = ["给定 α", "得到三个参数区间", "网格采样参数组合", "每组运行确定性 IDM", "逐时刻取 min/max", "得到 Fuzzy-IDM 包络"]
    for i, step in enumerate(steps):
        x = 1.2 + i * 5.2
        add_flow_box(s, x, 5.3, 4.3, 1.5, step, fill=PptRGB(242,246,250), size=12)
        if i < len(steps)-1:
            add_arrow(s, x+4.3, 6.05, x+5.2, 6.05)
    add_textbox(s, 3.0, 9.1, 27.0, 2.1,
                "Ωₐ = [a̲ₐ,āₐ] × [b̲ₐ,b̄ₐ] × [T̲ₐ,T̄ₐ]\nα=1 时三个区间退化为核心值，对应确定性 IDM。", 20, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)

    # 10
    s = prs.slides.add_slide(blank); add_title(s, "仿真设计"); add_footer(s, 10)
    table(s, 1.2, 2.35, 15.8, 7.7, [
        ["参数", "取值"],
        ["aₘₐₓ", "(0.8, 1.0, 1.2) m/s²"],
        ["a꜀ₒₘf", "(1.2, 1.5, 1.8) m/s²"],
        ["T", "(1.2, 1.5, 1.8) s"],
        ["Ṽₙ", "108 km/h"],
        ["Sⱼₐₘ / L / β", "2 m / 5 m / 4"],
        ["Δt / α", "0.05 s / 0,0.25,0.5,0.75,1"]
    ], col_widths=[5, 10.8], font_size=12)
    add_textbox(s, 18.2, 2.6, 12.8, 1.15, "三种典型交通场景", 18, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_bullets(s, 18.4, 4.3, 12.3, 5.7, [
        "匀加速前车：54 km/h 加速至 72 km/h 后保持恒速。",
        "前车减速停车：10 s 后以 1.5 m/s² 制动至停车。",
        "周期性扰动：60 s 稳定跟驰后叠加正弦速度扰动。"
    ], size=15)
    add_textbox(s, 18.4, 10.6, 12.3, 1.3,
                "每种场景均设置 5 辆跟驰车，用于观察响应滞后和扰动传播。", 14, PptRGB(30,30,30), fill=PptRGB(242,246,250), line=PptRGB(210,210,210), radius=True)

    # 11
    s = prs.slides.add_slide(blank); add_title(s, "匀加速场景结果"); add_footer(s, 11)
    add_picture_fit(s, IMG_DIR / "accelerating_velocity.png", 1.0, 2.15, 15.4, 8.7, "速度响应")
    add_picture_fit(s, IMG_DIR / "accelerating_acceleration.png", 17.0, 2.15, 15.4, 8.7, "加速度响应")
    add_bullets(s, 1.4, 11.6, 30.6, 1.9, [
        "跟驰车1响应更快，跟驰车5存在明显滞后；这是车队逐级响应造成的正常传播现象。",
        "Fuzzy-IDM 包络体现参数不确定性：α=0 包络最宽，α=0.5 更窄，接近稳定状态后包络收缩。"
    ], size=13.5)

    # 12
    s = prs.slides.add_slide(blank); add_title(s, "减速停车场景结果"); add_footer(s, 12)
    add_picture_fit(s, IMG_DIR / "braking_velocity.png", 1.0, 2.2, 10.2, 7.3, "速度")
    add_picture_fit(s, IMG_DIR / "braking_gap.png", 11.8, 2.2, 10.2, 7.3, "净间距")
    add_picture_fit(s, IMG_DIR / "braking_acceleration.png", 22.6, 2.2, 10.2, 7.3, "加速度")
    add_bullets(s, 1.4, 10.6, 30.6, 2.55, [
        "减速停车对参数更敏感，T、a꜀ₒₘf、aₘₐₓ 共同影响制动时机和制动强度。",
        "跟驰车5比跟驰车1更晚进入强制动，体现响应滞后。",
        "所有轨迹保持正净间距，当前参数范围内未出现碰撞。"
    ], size=13.5)

    # 13
    s = prs.slides.add_slide(blank); add_title(s, "周期扰动与交通流稳定性"); add_footer(s, 13)
    add_picture_fit(s, IMG_DIR / "periodic_velocity.png", 1.0, 2.15, 15.4, 8.7, "速度扰动传播")
    add_picture_fit(s, IMG_DIR / "periodic_stability_gain.png", 17.0, 2.15, 15.4, 8.7, "稳定性 G")
    add_bullets(s, 1.4, 11.6, 30.6, 1.9, [
        "Gᵢ=σ(Vᵢ)/σ(V₀)，G<1 表示速度扰动沿车队传播时衰减。",
        "当前结果中 G 随车辆序号下降，Fuzzy-IDM α=0 上界仍小于 1，说明该参数范围内速度扰动总体衰减。"
    ], size=13.5)

    # 14
    s = prs.slides.add_slide(blank); add_title(s, "跟驰行为波动性分析"); add_footer(s, 14)
    add_picture_fit(s, IMG_DIR / "periodic_acceleration.png", 1.0, 2.15, 15.4, 8.7, "周期扰动下加速度")
    add_picture_fit(s, IMG_DIR / "periodic_acceleration_fluctuation.png", 17.0, 2.15, 15.4, 8.7, "加速度标准差")
    add_bullets(s, 1.4, 11.6, 30.6, 1.9, [
        "加速度标准差用于衡量跟驰行为波动强度。",
        "确定性 IDM 加速度标准差沿车队下降；Fuzzy-IDM 给出波动范围，反映参数不确定性带来的预测离散程度。"
    ], size=13.5)

    # 15
    s = prs.slides.add_slide(blank); add_title(s, "阶段三构思：多模态输入下的模糊神经微分方程"); add_footer(s, 15)
    add_flow_box(s, 1.2, 2.6, 5.6, 1.2, "图像流", fill=PptRGB(242,246,250))
    add_flow_box(s, 1.2, 5.4, 5.6, 1.2, "点云流", fill=PptRGB(242,246,250))
    add_flow_box(s, 7.8, 2.6, 5.8, 1.2, "CNN/Transformer\n检测特征", fill=PptRGB(242,246,250), size=11)
    add_flow_box(s, 7.8, 5.4, 5.8, 1.2, "PointNet/Voxel/BEV\n空间特征", fill=PptRGB(242,246,250), size=11)
    add_flow_box(s, 14.8, 4.0, 5.7, 1.35, "模糊状态\nS̃ₙ,Ṽₙ,ΔṼₙ", fill=LIGHT_BLUE, size=12)
    add_flow_box(s, 21.3, 4.0, 5.7, 1.35, "FNN 模糊嵌入层\nũ(t) 或 θ̃(t)", fill=PptRGB(252,228,214), size=12)
    add_flow_box(s, 27.8, 4.0, 5.4, 1.35, "Fuzzy Neural ODE\nD_gH x̃/dt=fθ(x̃,ũ,t)", fill=PptRGB(226,239,218), size=11)
    for a in [(6.8,3.2,7.8,3.2),(6.8,6.0,7.8,6.0),(13.6,3.2,14.8,4.65),(13.6,6.0,14.8,4.65),(20.5,4.65,21.3,4.65),(27.0,4.65,27.8,4.65)]:
        add_arrow(s, *a)
    add_flow_box(s, 13.2, 8.5, 6.0, 1.2, "模糊数值求解", fill=PptRGB(242,246,250))
    add_flow_box(s, 21.0, 8.5, 6.0, 1.2, "轨迹预测/状态估计", fill=PptRGB(242,246,250))
    add_arrow(s, 30.5, 5.35, 16.2, 8.5); add_arrow(s, 19.2, 9.1, 21.0, 9.1)
    add_textbox(s, 1.4, 11.5, 31.0, 1.3,
                "本页是下一阶段构思：当前尚未完成多模态实验，后续需要通过文献调研、数据准备和简化模型验证逐步推进。", 14, PptRGB(90,55,0),
                align=PP_ALIGN.CENTER, fill=PptRGB(255,242,204), line=PptRGB(240,190,100), radius=True)

    # 16
    s = prs.slides.add_slide(blank); add_title(s, "总结与下一步"); add_footer(s, 16)
    add_textbox(s, 2.0, 2.6, 13.6, 1.0, "阶段性总结", 19, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_bullets(s, 2.2, 4.0, 13.0, 5.5, [
        "已完成模糊数、α-截集、H/gH 差和 FDE 数值求解的基础学习。",
        "已完成三参数 Fuzzy-IDM 建模理解与 Python 仿真。",
        "初步观察到参数不确定性会随跟驰动力学传播，并在周期扰动中体现为预测区间。"
    ], size=15)
    add_textbox(s, 18.0, 2.6, 13.6, 1.0, "下一步计划", 19, BLUE, True, PP_ALIGN.CENTER, fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_bullets(s, 18.2, 4.0, 13.0, 5.5, [
        "整理多模态感知不确定性、模糊神经网络和 Neural ODE 文献。",
        "完善阶段三构思文档，明确状态变量、输入变量和 α-截集求解方式。",
        "尝试设计 Multimodal Fuzzy Neural ODE 的模型框架、损失函数和初步实验流程。"
    ], size=15)
    add_textbox(s, 6.0, 11.2, 21.5, 1.05, "汇报完毕，恳请老师和同学批评指正", 20, BLUE, True, PP_ALIGN.CENTER, fill=PptRGB(242,246,250), line=BLUE, radius=True)

    prs.save(str(OUT_PPT))


def set_doc_style(doc):
    styles = doc.styles
    styles["Normal"].font.name = FONT_CN
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        if name in styles:
            styles[name].font.name = FONT_CN
            styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)


def add_doc_para(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = FONT_CN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        r.font.size = Pt(10.5)
        rest = text[len(bold_prefix):]
        r2 = p.add_run(rest)
        r2.font.name = FONT_CN
        r2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    else:
        r = p.add_run(text)
        r.font.name = FONT_CN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    p.paragraph_format.line_spacing = 1.25
    return p


def create_stage3_doc():
    doc = Document()
    set_doc_style(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0); sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.4)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("阶段三：多模态输入下的模糊神经微分方程构思")
    r.font.name = FONT_CN; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN); r.bold = True; r.font.size = Pt(18)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = sub.add_run("阶段性文档初稿")
    rr.font.name = FONT_CN; rr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN); rr.font.size = Pt(12)

    sections = [
        ("一、研究问题与动机", [
            "相机图像和激光雷达点云已经成为智能交通与自动驾驶场景中获取车辆状态的重要数据来源，但两类传感器都存在不可忽略的不确定性。相机容易受到光照、天气、阴影和遮挡影响，导致目标检测框、车辆边界、类别判断和速度估计出现不精确；激光雷达在远距离、雨雾或遮挡场景中点云变得稀疏，车辆轮廓、距离和相对速度估计也会出现模糊性。",
            "这些感知层面的不确定性会继续传递到微观交通模型中，使速度 Vₙ、实际净间距 Sₙ、相对速度 ΔVₙ 以及加速度 aₙ 的预测不再只是单点结果。阶段一和阶段二已经说明，模糊数与 α-截集可以用于描述非概率型不确定性，因此阶段三可以进一步探索：如何把多模态感知输入中的不确定性转化为模糊状态，并与神经微分方程结合，形成既能利用数据表达能力、又能输出预测区间的动力学模型。"
        ]),
        ("二、多模态输入层", [
            "图像流可以通过 CNN、Transformer 或目标检测网络提取车辆检测框、类别、中心位置、车道关系和速度估计等特征。检测框的置信度、边界抖动、遮挡比例和连续帧跟踪稳定性可作为不确定性来源，用于构造车辆位置和速度的模糊表示。",
            "点云流可以通过 PointNet、VoxelNet、BEV 表达或点云检测网络提取车辆三维中心、尺寸、相对距离、相对速度和空间结构特征。点云密度、远距离稀疏程度、遮挡点比例和检测置信度可用于表征距离与轮廓估计的不确定性。",
            "在融合层面，可以把检测置信度、边界扰动、点云稀疏程度映射为模糊状态，例如模糊间距 S̃ₙ、模糊速度 Ṽₙ、模糊相对速度 ΔṼₙ。这里的波浪线用于表示模糊状态；需要与经典 IDM 中 Ṽₙ、S̃ₙ、T̃ₙ 表示期望量的符号习惯区分开。"
        ]),
        ("三、模糊嵌入层", [
            "模糊嵌入层的作用是把多模态特征转化为可进入动力学模型的模糊状态或模糊参数。可以使用 Fuzzy Neural Network（FNN）对图像特征 z_img(t) 与点云特征 z_lidar(t) 进行映射：",
            "ũ(t)=FNNϕ(z_img(t),z_lidar(t))",
            "或者将多模态特征映射为动态模糊参数：",
            "θ̃(t)=FNNϕ(z_img(t),z_lidar(t))",
            "其中 θ̃ 可包括 aₘₐₓ、a꜀ₒₘf、T 等动态模糊参数。与阶段二中固定三角模糊参数不同，阶段三可以设想让参数随场景、驾驶行为和感知置信度动态变化。"
        ]),
        ("四、动力学层：Multimodal Fuzzy Neural ODE", [
            "阶段三拟构思的动力学层可写为：",
            "D_gH x̃(t)/dt = fθ(x̃(t),ũ(t),t)",
            "其中 x̃(t) 表示模糊交通状态，ũ(t) 表示多模态模糊输入，fθ 由神经网络参数化。gH 导数用于支持模糊区间随时间扩张或收缩的情形，使模型能够描述不确定性传播过程。",
            "该模型与传统模型的关系可以理解为：IDM 是人工设定的物理跟驰动力学；Fuzzy-IDM 是参数模糊化后的物理模型；Multimodal Fuzzy Neural ODE 则是在多模态感知输入下学习动力学函数，并尝试保留模糊不确定性的表达能力。"
        ]),
        ("五、求解与训练流程", [
            "初步算法流程可设计为：输入同步图像和点云；分别提取多模态特征；根据检测置信度、边界扰动和点云稀疏程度构造模糊状态或模糊参数；通过 α-截集把模糊动力学转化为一族区间 ODE；用数值求解器预测未来轨迹；根据真实轨迹计算损失；通过反向传播更新神经网络参数。",
            "损失函数可包括四类：中心轨迹误差，用于约束预测中心接近真实轨迹；上下界覆盖损失，用于鼓励真实轨迹落入预测区间；区间宽度正则项，用于避免预测区间无意义地过宽；物理约束损失，用于保证速度非负、间距非负以及加速度处于合理范围。"
        ]),
        ("六、可行性与难点", [
            "优势方面，该框架有望同时利用物理模型、模糊不确定性和神经网络表达能力，输出预测区间而不是单点预测，更适合多模态感知不确定场景。通过 α-截集表达预测区间，也便于与阶段二 Fuzzy-IDM 的方法延续起来。",
            "难点方面，gH 导数与神经网络耦合后数学表达会更加复杂；多 α 水平、多参数组合和神经 ODE 求解会带来较高计算成本；真实数据中很难直接获得模糊标签；图像与点云的同步、标定和误差传递建模要求较高；训练稳定性和物理约束损失的设计也需要进一步调研。"
        ]),
        ("七、小结", [
            "本阶段内容目前属于框架构思，不是已经完成的实验工作。后续需要先进行系统文献调研，明确模糊微分方程、神经 ODE、多模态感知不确定性和模糊神经网络之间的可衔接点；然后再进行模型简化、实验数据准备和初步代码验证。"
        ]),
    ]
    for heading, paras in sections:
        doc.add_heading(heading, level=1)
        for para in paras:
            add_doc_para(doc, para)

    doc.add_heading("附：初步框架关系表", level=1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    for i, text in enumerate(["模型层次", "核心思想", "阶段三中的作用"]):
        hdr[i].text = text
    for row in [
        ["IDM", "人工设定跟驰动力学", "提供物理先验和符号参照"],
        ["Fuzzy-IDM", "将关键驾驶参数模糊化", "提供 α-截集与包络求解思路"],
        ["Neural ODE", "用神经网络学习连续时间动力学", "提供数据驱动动力学表达"],
        ["Multimodal Fuzzy Neural ODE", "多模态输入 + 模糊状态 + 神经动力学", "阶段三拟构思的综合框架"]
    ]:
        cells = tbl.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    doc.save(str(OUT_STAGE3))


LITERATURE = [
    ["A 模糊微分方程理论", "Fuzzy differential equations", "模糊微分方程", "O. Kaleva", 1987, "Fuzzy Sets and Systems", "", "FDE 基础理论来源，说明模糊初值问题的基本形式。", "第4-6页；阶段三理论基础", "https://doi.org/10.1016/0165-0114(87)90029-7", ""],
    ["A 模糊微分方程理论", "Differentials of fuzzy functions", "模糊函数的微分", "M. L. Puri; D. A. Ralescu", 1983, "Journal of Mathematical Analysis and Applications", "", "Hukuhara 差与模糊函数微分的重要基础。", "第5页；阶段三动力学层", "https://doi.org/10.1016/0022-247X(83)90169-5", ""],
    ["A 模糊微分方程理论", "Generalizations of the differentiability of fuzzy-number-valued functions with applications to fuzzy differential equations", "模糊数值函数可微性的推广及其在模糊微分方程中的应用", "B. Bede; S. G. Gal", 2005, "Fuzzy Sets and Systems", "", "gH 可微性与 FDE 解的理论支撑。", "第5页；阶段三动力学层", "https://doi.org/10.1016/j.fss.2004.08.001", ""],
    ["A 模糊微分方程理论", "First order linear fuzzy differential equations under generalized differentiability", "广义可微性下的一阶线性模糊微分方程", "B. Bede; I. J. Rudas; A. L. Bencsik", 2007, "Information Sciences", "", "说明广义可微框架下 FDE 求解方式。", "第5-6页；阶段三求解流程", "https://doi.org/10.1016/j.ins.2007.02.021", ""],
    ["A 模糊微分方程理论", "Numerical methods for fuzzy differential equations under generalized differentiability", "广义可微性下模糊微分方程的数值方法", "T. Allahviranloo; S. Salahshour; S. Abbasbandy", 2012, "Applied Mathematical Modelling", "", "为 Euler/RK 与 α-截集数值求解提供依据。", "第6页；阶段三求解流程", "https://doi.org/10.1016/j.apm.2011.08.018", ""],
    ["B IDM及扩展模型", "Congested traffic states in empirical observations and microscopic simulations", "经验观测与微观仿真中的拥堵交通状态", "M. Treiber; A. Hennecke; D. Helbing", 2000, "Physical Review E", "Treiber_2000_Congested_traffic_states.pdf", "IDM 经典出处之一，用于解释跟驰模型和交通流现象。", "第7、10页", "https://doi.org/10.1103/PhysRevE.62.1805", "https://arxiv.org/pdf/cond-mat/0002177"],
    ["B IDM及扩展模型", "Enhanced intelligent driver model to access the impact of driving strategies on traffic capacity", "增强型智能驾驶员模型及驾驶策略对通行能力的影响", "A. Kesting; M. Treiber; D. Helbing", 2010, "Philosophical Transactions of the Royal Society A", "", "IDM 扩展与驾驶策略分析，可支撑模型背景。", "第7、13页", "https://doi.org/10.1098/rsta.2010.0084", ""],
    ["B IDM及扩展模型", "Mechanisms for spatio-temporal pattern formation in highway traffic models", "高速公路交通模型中时空模式形成机制", "R. E. Wilson", 2008, "Philosophical Transactions of the Royal Society A", "", "交通流扰动传播和稳定性理论参考。", "第13-14页", "https://doi.org/10.1098/rsta.2008.0018", ""],
    ["B IDM及扩展模型", "A car-following model based on fuzzy inference system", "基于模糊推理系统的跟驰模型", "S. Kikuchi; P. Chakroborty", 1992, "Transportation Research Record", "", "模糊方法进入跟驰行为建模的早期参考。", "第8页；阶段三文献调研", "https://trid.trb.org/view/370643", ""],
    ["B IDM及扩展模型", "Traffic Flow Dynamics: Data, Models and Simulation", "交通流动力学：数据、模型与仿真", "M. Treiber; A. Kesting", 2013, "Springer", "", "系统介绍 IDM、稳定性与仿真方法。", "第7、10、13页", "https://doi.org/10.1007/978-3-642-32460-4", ""],
    ["C Neural ODE", "Neural Ordinary Differential Equations", "神经常微分方程", "R. T. Q. Chen; Y. Rubanova; J. Bettencourt; D. Duvenaud", 2018, "NeurIPS", "Chen_2018_Neural_ODE.pdf", "阶段三 Fuzzy Neural ODE 的核心方法来源。", "第15页；阶段三文档", "https://arxiv.org/abs/1806.07366", "https://arxiv.org/pdf/1806.07366"],
    ["C Neural ODE", "Latent ODEs for Irregularly-Sampled Time Series", "面向非规则采样时间序列的潜在 ODE", "Y. Rubanova; R. T. Q. Chen; D. Duvenaud", 2019, "NeurIPS", "Rubanova_2019_Latent_ODE.pdf", "交通状态时间序列不规则采样与连续动力学建模参考。", "阶段三求解与训练", "https://arxiv.org/abs/1907.03907", "https://arxiv.org/pdf/1907.03907"],
    ["C Neural ODE", "Augmented Neural ODEs", "增广神经 ODE", "E. Dupont; A. Doucet; Y. W. Teh", 2019, "NeurIPS", "Dupont_2019_Augmented_Neural_ODEs.pdf", "用于理解普通 Neural ODE 表达能力不足时的扩展方式。", "阶段三模型简化与扩展", "https://arxiv.org/abs/1904.01681", "https://arxiv.org/pdf/1904.01681"],
    ["C Neural ODE", "Physics-informed neural networks: A deep learning framework for solving forward and inverse problems involving nonlinear partial differential equations", "物理信息神经网络", "M. Raissi; P. Perdikaris; G. E. Karniadakis", 2019, "Journal of Computational Physics", "Raissi_2019_PINNs.pdf", "为阶段三物理约束损失提供借鉴。", "阶段三损失函数", "https://doi.org/10.1016/j.jcp.2018.10.045", "https://arxiv.org/pdf/1711.10561"],
    ["D 多模态感知不确定性建模", "What Uncertainties Do We Need in Bayesian Deep Learning for Computer Vision?", "计算机视觉中的贝叶斯深度学习需要哪些不确定性", "A. Kendall; Y. Gal", 2017, "NeurIPS", "Kendall_2017_Bayesian_Uncertainty.pdf", "用于区分感知中的认知不确定性和数据噪声。", "第2、15页；阶段三动机", "https://arxiv.org/abs/1703.04977", "https://arxiv.org/pdf/1703.04977"],
    ["D 多模态感知不确定性建模", "VoxelNet: End-to-End Learning for Point Cloud Based 3D Object Detection", "VoxelNet：基于点云的端到端三维目标检测", "Y. Zhou; O. Tuzel", 2018, "CVPR", "Zhou_2018_VoxelNet.pdf", "点云流特征提取和 3D 检测基础。", "阶段三多模态输入层", "https://arxiv.org/abs/1711.06396", "https://arxiv.org/pdf/1711.06396"],
    ["D 多模态感知不确定性建模", "PointPillars: Fast Encoders for Object Detection from Point Clouds", "PointPillars：点云目标检测快速编码器", "A. H. Lang; S. Vora; H. Caesar; L. Zhou; J. Yang; O. Beijbom", 2019, "CVPR", "Lang_2019_PointPillars.pdf", "BEV/点云编码方式参考。", "阶段三多模态输入层", "https://arxiv.org/abs/1812.05784", "https://arxiv.org/pdf/1812.05784"],
    ["D 多模态感知不确定性建模", "Frustum PointNets for 3D Object Detection from RGB-D Data", "基于 RGB-D 数据的视锥 PointNets 三维目标检测", "C. R. Qi; W. Liu; C. Wu; H. Su; L. J. Guibas", 2018, "CVPR", "Qi_2018_Frustum_PointNets.pdf", "图像与点云几何融合的代表性方法。", "阶段三多模态输入层", "https://arxiv.org/abs/1711.08488", "https://arxiv.org/pdf/1711.08488"],
    ["D 多模态感知不确定性建模", "Deep Continuous Fusion for Multi-Sensor 3D Object Detection", "用于多传感器三维目标检测的深度连续融合", "M. Liang; B. Yang; S. Wang; R. Urtasun", 2018, "ECCV", "Liang_2018_Deep_Continuous_Fusion.pdf", "相机-激光雷达融合结构参考。", "阶段三多模态输入层", "https://arxiv.org/abs/1805.00772", "https://arxiv.org/pdf/1805.00772"],
    ["D 多模态感知不确定性建模", "LaserNet: An Efficient Probabilistic 3D Object Detector for Autonomous Driving", "LaserNet：面向自动驾驶的高效概率三维目标检测器", "G. Meyer; A. Laddha; E. Kee; C. Vallespi-Gonzalez; C. K. Wellington", 2019, "CVPR", "Meyer_2019_LaserNet.pdf", "概率式 3D 检测输出可启发模糊状态构造。", "阶段三不确定性建模", "https://arxiv.org/abs/1903.08701", "https://arxiv.org/pdf/1903.08701"],
    ["D 多模态感知不确定性建模", "nuScenes: A multimodal dataset for autonomous driving", "nuScenes：自动驾驶多模态数据集", "H. Caesar et al.", 2020, "CVPR", "Caesar_2020_nuScenes.pdf", "可作为后续多模态实验数据参考。", "阶段三数据准备", "https://arxiv.org/abs/1903.11027", "https://arxiv.org/pdf/1903.11027"],
    ["E 模糊神经网络在交通中的应用", "ANFIS: Adaptive-Network-Based Fuzzy Inference System", "ANFIS：基于自适应网络的模糊推理系统", "J.-S. R. Jang", 1993, "IEEE Transactions on Systems, Man, and Cybernetics", "", "模糊神经网络的经典基础，可支撑 FNN 嵌入层。", "第15页；阶段三模糊嵌入层", "https://doi.org/10.1109/21.256541", ""],
    ["E 模糊神经网络在交通中的应用", "Short-term traffic forecasting: Where we are and where we’re going", "短时交通预测：现状与发展方向", "E. I. Vlahogianni; M. G. Karlaftis; J. C. Golias", 2014, "Transportation Research Part C", "", "交通预测任务综述，可用于阶段三问题定位。", "阶段三研究问题", "https://doi.org/10.1016/j.trc.2014.01.005", ""],
    ["B IDM及扩展模型", "智能驾驶员模型及稳定性分析", "Intelligent driver model and stability analysis", "陈喜群; 杨新苗; 史其信; 秦旭彦", 2008, "系统仿真学报", "", "中文 IDM 与稳定性分析参考。", "第7、13页", "https://kns.cnki.net/kcms/detail/detail.aspx?dbcode=CJFD&filename=XTLY200808049", ""],
    ["B IDM及扩展模型", "考虑后视效应和速度差的智能驾驶模型仿真", "Simulation of intelligent driver model considering backward-looking effect and velocity difference", "陈然; 薛郁", 2013, "物理学报", "", "中文跟驰模型扩展与扰动传播参考。", "第13页", "https://wulixb.iphy.ac.cn/article/id/54817", ""],
    ["A 模糊微分方程理论", "半线性时滞模糊微分方程解的存在性", "Existence of solutions for semilinear delay fuzzy differential equations", "刘娟; 肖建中; 姚忠豪", 2006, "工程数学学报", "", "中文 FDE 解存在性参考。", "第5页；阶段三理论基础", "https://www.gcsxjsgc.com/CN/Y2006/V23/I5/879", ""],
    ["E 模糊神经网络在交通中的应用", "基于自适应神经模糊推理系统的跟驰模型研究", "Car-following model based on adaptive neuro-fuzzy inference system", "陈阳伍; 张一斌; 张阳", 2009, "公路交通技术", "", "ANFIS 与跟驰行为建模的中文应用参考。", "阶段三模糊嵌入层", "https://kns.cnki.net/kns8/defaultresult/index?kw=%E5%9F%BA%E4%BA%8E%E8%87%AA%E9%80%82%E5%BA%94%E7%A5%9E%E7%BB%8F%E6%A8%A1%E7%B3%8A%E6%8E%A8%E7%90%86%E7%B3%BB%E7%BB%9F%E7%9A%84%E8%B7%9F%E9%A9%B0%E6%A8%A1%E5%9E%8B%E7%A0%94%E7%A9%B6", ""],
    ["E 模糊神经网络在交通中的应用", "基于模糊神经网络的城市高速公路入口匝道控制算法", "Freeway ramp control algorithm based on neurofuzzy networks", "陈德望; 王飞跃; 陈龙", 2003, "交通运输工程学报", "", "模糊神经网络在交通控制中的中文应用，可辅助说明交通领域可行性。", "阶段三可行性与应用背景", "https://transport.chd.edu.cn/article/id/200302023?viewType=HTML", ""],
]


def safe_filename(name):
    name = re.sub(r"[\\/:*?\"<>|]+", "_", name)
    return name[:120]


def download_public_pdfs():
    PDF_DIR.mkdir(exist_ok=True)
    headers = {"User-Agent": "Mozilla/5.0"}
    ok = 0
    fail = []
    for rec in LITERATURE:
        filename = rec[6]
        url = rec[10]
        if not filename or not url:
            continue
        out = PDF_DIR / safe_filename(filename)
        if out.exists() and out.stat().st_size > 50000:
            ok += 1
            continue
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=45) as resp:
                data = resp.read()
            if data[:4] == b"%PDF" or len(data) > 50000:
                out.write_bytes(data)
                ok += 1
            else:
                fail.append((filename, "下载内容不是有效 PDF"))
        except Exception as e:
            fail.append((filename, str(e)))
    return ok, fail


def create_literature_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "文献清单"
    headers = ["中文分类", "英文题名", "中文译名", "作者", "年份", "来源/期刊", "PDF文件名", "与本课题的关系", "可用于报告哪一部分", "可访问链接", "公开PDF链接"]
    ws.append(headers)
    for rec in LITERATURE:
        row = list(rec)
        if not row[6]:
            row[6] = "未找到公开PDF"
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    border = Border(left=Side(style="thin", color="D9E2F3"), right=Side(style="thin", color="D9E2F3"),
                    top=Side(style="thin", color="D9E2F3"), bottom=Side(style="thin", color="D9E2F3"))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.font = Font(name="等线", size=10)
    for cell in ws[1]:
        cell.font = Font(name="等线", bold=True, size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [16, 36, 30, 28, 8, 28, 28, 38, 28, 45, 45]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(str(OUT_LIT))


def create_lit_review_doc(download_ok, download_fail):
    doc = Document()
    set_doc_style(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("阶段三文献调研小结")
    r.font.name = FONT_CN; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN); r.bold = True; r.font.size = Pt(18)
    english_count = sum(1 for rec in LITERATURE if isinstance(rec[1], str) and all(ord(ch) < 128 for ch in rec[1]))
    chinese_count = len(LITERATURE) - english_count
    add_doc_para(doc, f"本次围绕阶段三“多模态输入下的模糊神经微分方程构思”整理了 {len(LITERATURE)} 篇可追溯文献，其中英文文献 {english_count} 篇，中文文献 {chinese_count} 篇。已尝试下载公开可访问 PDF，成功下载 {download_ok} 篇；其余文献在清单中保留 DOI、期刊页面、CNKI/期刊页面或检索入口，并标注“未找到公开PDF”。")
    topics = [
        ("一、模糊微分方程理论", "Kaleva、Puri 与 Ralescu、Bede 与 Gal 等文献构成 FDE 与 gH 可微性的理论基础。它们支撑阶段报告中 α-截集、H/gH 差和模糊初值问题的表述，也为阶段三中 D_gH x̃(t)/dt 的写法提供依据。"),
        ("二、IDM 与交通流稳定性", "Treiber、Kesting、Wilson 等文献可用于解释 IDM 的物理意义、车队扰动传播和稳定性指标。中文文献《智能驾驶员模型及稳定性分析》和相关跟驰模型扩展文献可作为组会汇报中的中文背景参考。"),
        ("三、Neural ODE 与物理约束学习", "Neural ODE、Latent ODE、Augmented Neural ODE 与 PINNs 文献为阶段三提供连续时间神经动力学、非规则时间序列建模、表达能力扩展和物理约束损失函数设计思路。"),
        ("四、多模态感知不确定性", "Kendall 与 Gal 的不确定性分类，以及 VoxelNet、PointPillars、Frustum PointNets、Deep Continuous Fusion、LaserNet 和 nuScenes 等文献，可用于说明图像与点云输入、检测置信度、点云稀疏性和多模态融合如何产生并表达不确定性。"),
        ("五、模糊神经网络在交通中的应用", "ANFIS 与中文交通应用文献可支撑 FNN 模糊嵌入层的构思。后续需要进一步筛选与交通流预测、跟驰行为建模更直接相关的中文全文。"),
    ]
    for heading, text in topics:
        doc.add_heading(heading, level=1)
        add_doc_para(doc, text)
    doc.add_heading("后续建议", level=1)
    for item in [
        "优先精读 FDE/gH 可微性、Neural ODE、感知不确定性三类英文核心文献。",
        "中文文献主要用于补充国内 IDM、交通流稳定性和模糊神经网络应用背景。",
        "阶段三初稿中暂不急于追求完整复杂模型，可先设计“多模态特征 → 模糊参数 → Fuzzy-IDM/Neural ODE”的简化版本。",
        "如果后续需要正式论文综述，建议继续使用学校数据库补齐中文文献全文和引用格式。"
    ]:
        add_doc_para(doc, "• " + item)
    if download_fail:
        doc.add_heading("公开 PDF 下载说明", level=1)
        add_doc_para(doc, "以下公开 PDF 链接在本次自动下载中未成功，已在 Excel 中保留链接，后续可手动复查：")
        for filename, reason in download_fail[:12]:
            add_doc_para(doc, f"• {filename}：{reason}")
    doc.save(str(OUT_REVIEW))


def create_all():
    create_ppt()
    create_stage3_doc()
    ok, fail = download_public_pdfs()
    create_literature_xlsx()
    create_lit_review_doc(ok, fail)
    print("created", OUT_PPT)
    print("created", OUT_STAGE3)
    print("created", OUT_LIT)
    print("created", OUT_REVIEW)
    print("pdf_download_ok", ok)
    if fail:
        print("pdf_download_fail", len(fail))
        for item in fail[:20]:
            print(item[0], item[1])


if __name__ == "__main__":
    create_all()
